"""AUTH-11 — le cliquet des sorties d'une collection : jouée et propre, ou DÉCLARÉE.

`test_depot_export._textes_de_depot` se présentait comme « tout ce qu'une collection fait
sortir » et en regardait quatre sorties sur une vingtaine : ni l'archive, ni le classeur,
ni la fiche de description, ni la voie ShareDocs, ni aucune ligne de commande. Rien ne
fuyait en silence le jour où on l'a mesuré (2026-09-23) — mais c'est le mode d'échec que
`CLAUDE.md` impute à `test_csp` : une liste écrite à la main **oublie ce qu'on ajoute**
au lieu de perdre ce qu'elle voyait. Une sortie ajoutée demain n'y serait jamais entrée.

Le patron est celui de `test_autorisation.py` et de `test_sorties_identite.py` : ÉNUMÉRER
depuis une source que personne ne tient à la main, exiger que chaque surface ait été
TRANCHÉE, échouer sur celle qui ne figure nulle part — et sur la déclaration qui ne
désigne plus rien, parce qu'une liste périmée rassure.

**Ce qui est énuméré, et contre quoi.**

- Les ROUTES viennent de `inventaire_routes` (ARCH-2), plancher compris. Est une sortie
  toute route qui produit un fichier (le critère de DROIT-2, `_sort_un_fichier`, repris
  et non recopié) ou qui ENVOIE hors de l'instance (`sharedocs.upload`). Une route
  `/depot/` est jouée d'office, dans chaque format que son propre contrat déclare (le
  `pattern` de son paramètre `format`, confronté à `routes.depot._FORMATS`) et dans
  chaque position de ses autres paramètres ; un paramètre que ce fichier ne sait pas
  jouer fait échouer.
- Les OUTILS viennent de `tools/*.py`, et leurs entrées de leur propre `argparse`, lu
  par AST : options longues ET courtes, arguments positionnels, sous-commandes, chacune
  rangée sous la sous-commande qui la porte. Pour un outil JOUÉ, toute entrée est
  classée ; pour un outil HORS du contrôle, l'ensemble de ses entrées est FIGÉ — une
  option neuve y fait échouer aussi, parce que c'est par une option ajoutée qu'un outil
  qui « n'exporte rien » se met à exporter. Ce que l'AST ne tient pas est écrit avec
  `_entrees_argparse` : trois outils lisent `sys.argv` à la main, et une sortie qui ne
  dépend d'AUCUNE entrée (un fichier écrit d'office) n'est vue par personne.

**Ce que « propre » veut dire** — la définition du contrôle existant, sans l'affaiblir :
aucun terme d'une collection qu'on ne dépose pas (les sentinelles de
`vocabulaire_cloisonne`, catalogues ET liaisons), aucune charge du journal (les
marqueurs de `_semer_une_charge_par_cible`, plus les charges RÉELLES que le décor
produit en passant par l'API). S'y ajoutent le périmètre des ALBUMS — le titre de
l'album de Bravo, le texte et la note de sa région — et le marqueur de portée d'un run.
« Propre » veut donc dire « aucune de CES sentinelles », pas « aucune donnée de Bravo » :
un identifiant numérique, un compteur trop grand, une colonne que le décor ne remplit
pas passeraient. Et, dans l'autre sens, l'anti-vacuité du contrôle
existant devenue une ÉGALITÉ : chaque artefact porte exactement le vocabulaire que sa
famille annonce. C'est elle qui fait tomber un lecteur vacant — un classeur compressé
lu en octets ne « porte » rien, et paraîtrait propre.

**Ce qu'il ne fait pas, et qu'il ne faut pas lui prêter.** Il ferme la porte de l'OUBLI,
pas celle de l'erreur : une sortie dont le contenu fuit autre chose que ces sentinelles
passe. Son critère de « sortie » est celui de DROIT-2, et il en hérite l'angle mort : une
route qui rendrait le contenu d'une collection en JSON ordinaire, sans pièce jointe ni
`/export` ni `/depot/` dans son chemin, n'est une sortie pour aucun des deux (aucune
route n'est dans ce cas au 2026-10-07, vérifié par la relecture). Sa sentinelle est
ADMINISTRATRICE, donc il mesure l'exposition maximale et ne dit rien de ce qu'un compte restreint obtient (c'est `test_droit_export`). La voie ShareDocs
est jouée contre une DOUBLURE de l'envoi WebDAV : ce qui part réellement chez Huma-Num
n'est regardé par aucun test. Et le mode « corpus entier » des outils (sans
`--collection`) n'est pas jugé ici : il n'a pas de frontière à tenir.
"""
import ast
import importlib.util
import inspect
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import zipfile
from collections import namedtuple
from pathlib import Path

import pytest

import inventaire_routes
import main
from conftest import ADMIN, make_png
from test_depot_export import vocabulaire_cloisonne  # noqa: F401  (fixture importée)
from test_droit_export import _sort_un_fichier
from test_provenance_audit import _semer_une_charge_par_cible

REPO_ROOT = Path(__file__).resolve().parent.parent
TOOLS = REPO_ROOT / "tools"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

# Le périmètre d'un run tel que `activite.portee` le range — `{"planche_id": N}`, parfois
# `album_id`. Un identifiant numérique ne se cherche pas dans un document ; le marqueur
# est donc posé À CÔTÉ de lui, dans la même colonne : ce qu'on mesure est « la colonne
# part-elle », exactement comme les marqueurs `CHARGE-…` pour `avant`/`apres`.
PORTEE_AILLEURS = "PORTEE-BRAVO-9500"

# Le périmètre des ALBUMS. Les deux titres sont ceux de `vocabulaire_cloisonne` ; le
# texte et la note sont posés par `seme` sur une bulle de l'album de Bravo. « Corpus
# Bravo », le NOM de la collection, ne peut pas être un interdit : `gerer_collections.py
# lister` le rend légitimement — il liste les collections de l'instance.
ALBUM_ALPHA = "ALBUM-ALPHA-9001"
ALBUM_AILLEURS = "ALBUM-BRAVO-9002"
TEXTE_AILLEURS = "TEXTE-BRAVO-9600"
NOTE_AILLEURS = "NOTE-BRAVO-9601"

DEPOT = "/depot/"
DEPOSER = "/api/collections/{collection_id}/depot/deposer"
ADRESSE_IMAGES = "https://images.example.org/iiif"


# --------------------------------------------------------------------------- #
# Lire un artefact, quel que soit son emballage
# --------------------------------------------------------------------------- #
def _parties(blob: bytes, nom: str = "") -> dict:
    """Le contenu FOUILLABLE d'un artefact, partie par partie : {nom de partie: texte}.

    Un classeur XLSX est un zip dont les chaînes vivent, compressées, dans
    `sharedStrings.xml` : le chercher en octets ne voit RIEN, et le contrôle serait vert
    pour la mauvaise raison (la faute relevée le 2026-08-31 sur AUTH-5, puis le
    2026-09-24 ici même). Il est donc OUVERT, feuille par feuille, cellule par cellule ;
    une archive est dépliée membre par membre.

    Rendre des PARTIES et non un texte unique n'est pas un raffinement : une limite
    déclarée sur l'onglet `fiche` ne doit pas couvrir la même fuite dans l'onglet
    `vocabulaire`. Un écart se déclare à l'endroit où il est, pas sur le fichier entier.

    **Une partie qu'on ne sait pas lire fait ÉCHOUER, elle ne devient pas du bruit.** Un
    membre décodé « au mieux » (`errors="replace"`) se lit toujours — et un
    `tags_complet.txt.gz` glissé dans l'archive passait ainsi pour propre, ses sentinelles
    compressées. D'où deux exigences : une extension CONNUE comme texte, décodée
    STRICTEMENT ; ou une extension déclarée binaire dans `BINAIRES`, avec ce qu'on
    renonce à y lire.
    """
    if blob[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            membres = z.namelist()
            if "[Content_Types].xml" in membres:
                return _feuilles(blob, nom)
            parties = {}
            for m in membres:
                parties.update(_parties(z.read(m), f"{nom}{'/' if nom else ''}{m}"))
            return parties
    ext = Path(nom).suffix.lower()
    if ext in BINAIRES:
        return {nom: ""}
    assert ext in TEXTES, (
        f"« {nom} » : extension {ext!r} inconnue du lecteur. Si c'est du texte, l'ajouter "
        "à TEXTES ; si c'est un binaire, le déclarer dans BINAIRES avec sa raison — un "
        "contenu qu'on ne lit pas ne peut pas être dit propre.")
    try:
        return {nom: blob.decode("utf-8")}
    except UnicodeDecodeError as exc:
        raise AssertionError(
            f"« {nom} » ne se décode pas en UTF-8 ({exc}) : ce n'est pas le texte que "
            "son extension annonce, et le fouiller quand même ne prouverait rien.")


# Ce que le lecteur sait décoder. `""` est l'artefact nu — le corps d'une réponse, la
# sortie standard d'un outil — et il est décodé strictement comme les autres.
TEXTES = {"", ".json", ".jsonld", ".csv", ".txt", ".xml"}
# Ce qu'il renonce à lire, et pourquoi.
BINAIRES = {
    ".png": "le crop d'une figure (`POST /api/figures`) : des pixels. Ses blocs de texte "
            "PNG (`tEXt`) ne sont pas lus — Pillow n'en écrit aucun ici, et c'est une "
            "limite, pas une garantie.",
}


def _feuilles(blob: bytes, nom: str) -> dict:
    """Les feuilles d'un classeur, lues par `openpyxl` — la bibliothèque qui l'a écrit.

    Lu : le titre de chaque feuille, la valeur de chaque cellule (une formule sort en
    clair, lien compris) et son commentaire ; plus les PROPRIÉTÉS du classeur (titre,
    sujet, auteur, mots-clés, description…), rendues comme une partie à part.
    **Non lu, et c'est une limite** : les noms définis, les en-têtes et pieds de page,
    les validations de données, les images et graphiques, et tout XML que `openpyxl` ne
    modélise pas. `_ecrire_xlsx` n'en écrit aucun aujourd'hui.
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob))
    prefixe = f"{nom}/" if nom else ""
    parties = {}
    for ws in wb.worksheets:
        morceaux = [ws.title]
        for ligne in ws.iter_rows():
            for c in ligne:
                if c.value is not None:
                    morceaux.append(str(c.value))
                if c.comment is not None:
                    morceaux.append(str(c.comment.text))
        parties[prefixe + ws.title] = "\n".join(morceaux)
    p = wb.properties
    parties[prefixe + "(propriétés)"] = "\n".join(
        str(v) for v in (p.title, p.subject, p.creator, p.keywords, p.description,
                         p.category, p.lastModifiedBy, p.contentStatus, p.identifier)
        if v)
    return parties


def _brut(blob: bytes) -> str:
    """Les OCTETS d'une archive, membre par membre, sans rien décoder — pour la seule
    sauvegarde, dont le membre est un fichier SQLite : ses pages portent le texte en
    clair, et c'est exactement ce qu'on y cherche (AUTH-5 fait de même)."""
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        return "\n".join(z.read(m).decode("latin-1") for m in z.namelist())


def _fichiers(chemin: Path) -> dict:
    """Les parties de ce qu'un outil a écrit sur disque (un fichier, ou un dossier)."""
    if chemin.is_file():
        return _parties(chemin.read_bytes(), chemin.name)
    parties = {}
    for f in sorted(p for p in chemin.glob("**/*") if p.is_file()):
        parties.update(_parties(f.read_bytes(), f.relative_to(chemin).as_posix()))
    return parties


def _xlsx_lisible() -> bool:
    return importlib.util.find_spec("openpyxl") is not None


# --------------------------------------------------------------------------- #
# Le décor
# --------------------------------------------------------------------------- #
@pytest.fixture
def seme(client, db_path, data_dir, vocabulaire_cloisonne):
    """Le décor du contrôle existant, plus ce que le journal range.

    Rien n'est réinventé : les termes sont ceux de `vocabulaire_cloisonne` (trois portées,
    deux liaisons qui croisent la frontière), les charges celles de
    `_semer_une_charge_par_cible`, qui se relit lui-même en base. S'y ajoutent une
    activité dont le périmètre désigne l'album d'à côté, et le régime `public` des deux
    collections — sans lui le manifeste IIIF refuse `verbatim` (403) et sort sans images,
    c'est-à-dire qu'on mesurerait la sortie la plus PAUVRE. Comme la sentinelle
    administratrice d'AUTH-5, le décor vise l'exposition maximale.

    **Et l'album de Bravo reçoit un CONTENU** : une planche, une bulle, son texte, sa note
    et son tag. `vocabulaire_cloisonne` le laisse vide, si bien que le périmètre des
    ALBUMS n'était gardé par rien — un outil qui oubliait `--collection` pour ses albums
    sortait ceux de Bravo, et aucune sentinelle n'y était (mutant de la relecture,
    2026-10-07). Le titre, le texte et la note entrent dans les interdits.
    """
    v = vocabulaire_cloisonne
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        album = {r["titre"]: r["id"] for r in conn.execute("SELECT id, titre FROM albums")}
        region = conn.execute(
            "SELECT r.id FROM regions r JOIN planches p ON p.id = r.planche_id "
            "WHERE p.album_id = ?", (album[ALBUM_ALPHA],)).fetchone()["id"]
    finally:
        conn.close()

    pl = client.post(f"/api/albums/{album[ALBUM_AILLEURS]}/import", headers=ADMIN,
                     files={"file": ("b.png", make_png(), "image/png")})
    assert pl.status_code in (200, 201), pl.text
    reg = client.post(f"/api/planches/{pl.json()['id']}/regions", headers=ADMIN,
                      json={"type": "bulle", "x": 0, "y": 0, "w": 9, "h": 9})
    assert reg.status_code in (200, 201), reg.text
    region_bravo = reg.json()["id"]
    r = client.put(f"/api/regions/{region_bravo}", headers=ADMIN,
                   json={"ocr_texte": TEXTE_AILLEURS})
    assert r.status_code == 200, r.text
    r = client.put(f"/api/regions/{region_bravo}/annotation", headers=ADMIN,
                   json={"note": NOTE_AILLEURS, "tags": ["tag-bravo-9402"]})
    assert r.status_code == 200, r.text

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        cibles = _semer_une_charge_par_cible(conn)
        conn.execute("UPDATE collection SET statut_diffusion = 'public' WHERE id IN (?, ?)",
                     (v["alpha"]["id"], v["bravo"]["id"]))
        conn.execute(
            "INSERT INTO activite (type, agent, agent_type, portee) "
            "VALUES ('ocr', 'easyocr', 'moteur', ?)",
            (json.dumps({"album_id": album[ALBUM_AILLEURS],
                         "marqueur": PORTEE_AILLEURS}),))
        conn.commit()
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE)")       # les outils lisent à part
    finally:
        conn.close()
    interdits = (v["b"] + [f"CHARGE-{t}" for t in cibles]
                 + [PORTEE_AILLEURS, ALBUM_AILLEURS, TEXTE_AILLEURS, NOTE_AILLEURS])
    # Les interdits restent LIÉS au semis : le jour où un écart sera fermé et sa
    # sentinelle retirée d'une déclaration, elle ne doit pas pouvoir quitter cette liste
    # du même geste.
    assert set(v["b"]) <= set(interdits)
    return {**v, "db": db_path, "data": data_dir, "region_id": region,
            "region_bravo": region_bravo,
            "album_id": album[ALBUM_ALPHA],
            "charges": [f"CHARGE-{t}" for t in cibles],
            "propre": v["global"] + v["a"],
            "interdits": interdits}


# --------------------------------------------------------------------------- #
# Ce que chaque sortie JOUÉE doit porter — et rien d'autre
# --------------------------------------------------------------------------- #
# `porte` : les familles de termes que l'artefact publie, par préfixe de sentinelle. Le
# cliquet en fait une ÉGALITÉ avec ce qu'il lit — c'est l'anti-vacuité du contrôle
# existant (« le vocabulaire global ET celui d'Alpha doivent SORTIR »), et elle sert deux
# fois : un filtre qui emporterait tout tombe, et un LECTEUR qui ne verrait rien aussi.
# `temoin` : pour un artefact qui ne porte aucun terme, la chaîne qui prouve qu'on l'a lu.
Attendu = namedtuple("Attendu", "porte temoin raison")

TOUT = ("dom-", "dim-", "val", "tag-")
AXES = ("dom-", "dim-", "val")          # la fiche nomme les axes, elle COMPTE les tags
SUJETS = ("val", "tag-")                # la notice de dépôt : des sujets, pas des axes
TAGS = ("tag-",)                        # le manifeste : les tags posés sur les régions
AUCUN = ()

_DESC = "/api/collections/{collection_id}/depot/description"
_META = "/api/collections/{collection_id}/depot/metadonnees"
_IIIF = "/api/collections/{collection_id}/depot/iiif"

ATTENDUS = {
    # ---- Routes de dépôt -------------------------------------------------- #
    ("route", _DESC, "json"): Attendu(
        AXES, None, "le roll-up nomme domaines, dimensions et valeurs ; il compte les tags"),
    ("route", _DESC, "csv"): Attendu(
        AUCUN, "Corpus Alpha", "le catalogue : une ligne par ÉLÉMENT de métadonnée, des "
        "agrégats, aucun terme nommé"),
    ("route", _META, "json"): Attendu(TOUT, None, "l'arbre des enregistrements"),
    ("route", _META, "zip"): Attendu(TOUT, None, "les tables CSV, une par niveau"),
    ("route", _META, "xlsx"): Attendu(TOUT, None, "le classeur : tables, arbre et fiche"),
    ("route", _IIIF, "zip"): Attendu(
        TAGS, None, "collection + un manifeste par album : les tags voyagent en "
        "annotations `tagging`, les attributs n'y sont pas"),
    # ---- Autres routes jouées --------------------------------------------- #
    # L'export d'un album, au titre de la collection : ses tags sont ceux de CETTE
    # collection et les globaux, quel que soit qui exporte (AUTH-11, tranché le
    # 2026-10-07 — `socle._vocabulaire_d_export`). L'écart « À TRANCHER » que ce balayage
    # avait trouvé et déclaré ici est fermé : ces trois sorties n'ont plus d'entrée dans
    # ECARTS, et le tag de Bravo y est redevenu une fuite comme une autre. Le compte par
    # compte est joué à part, par `test_depot_export` (§ 2 ter).
    ("route", "/api/export/json", "json"): Attendu(TAGS, None, "l'export d'un album"),
    ("route", "/api/export/csv", "csv"): Attendu(TAGS, None, "l'export d'un album"),
    ("route", "/api/export/tei", "tei"): Attendu(TAGS, None, "l'export d'un album"),
    ("route", "/api/figures", "zip"): Attendu(
        AUCUN, "ALBUM-ALPHA-9001", "crop + légende + notice : une liste FERMÉE de "
        "mentions (`figure.CHAMPS`), aucune ne lit le vocabulaire"),
    # ---- Outils ----------------------------------------------------------- #
    ("outil", "description_collection.py", "--json"): Attendu(AXES, None, "= la route"),
    ("outil", "description_collection.py", "--csv"): Attendu(
        AUCUN, "Corpus Alpha", "= la route"),
    ("outil", "metadonnees_collection.py", "--json"): Attendu(TOUT, None, "= la route"),
    ("outil", "metadonnees_collection.py", "--csv-dir"): Attendu(
        TOUT, None, "les tables de l'archive, écrites à plat dans un dossier"),
    ("outil", "metadonnees_collection.py", "--zip"): Attendu(TOUT, None, "= la route"),
    ("outil", "metadonnees_collection.py", "--xlsx"): Attendu(TOUT, None, "= la route"),
    ("outil", "iiif_manifest.py", "--out-dir"): Attendu(TAGS, None, "= la route"),
    ("outil", "iiif_manifest.py", "(stdout)"): Attendu(
        TAGS, None, "sans `--out-dir` : le manifeste du PREMIER album, sur la sortie"),
    ("outil", "crosswalk_depot.py", "--out-dir"): Attendu(
        SUJETS, None, "notices Dublin Core et DataCite (JSON + XML)"),
    ("outil", "crosswalk_depot.py", "(stdout)"): Attendu(SUJETS, None, "le même, en JSON"),
    ("outil", "provenance_export.py", "--out-dir"): Attendu(
        AUCUN, "annotateur-", "PROV-JSON + TEI revisionDesc : des ACTES, sans charges"),
    ("outil", "provenance_export.py", "(stdout)"): Attendu(
        AUCUN, "annotateur-", "le même, en un document"),
    ("outil", "gerer_collections.py", "lister"): Attendu(
        AUCUN, "Corpus Alpha", "compte rendu de terminal : les collections"),
    ("outil", "gerer_collections.py", "montrer"): Attendu(
        AUCUN, "ALBUM-ALPHA-9001", "compte rendu de terminal : la fiche d'une collection"),
}


# --------------------------------------------------------------------------- #
# Les ÉCARTS déclarés — ce qu'une sortie jouée laisse passer, et pourquoi
# --------------------------------------------------------------------------- #
# Une entrée = une sortie → ({PARTIE : les sentinelles qui y passent}, la raison). Mesuré
# dans les deux sens ET partie par partie : une sentinelle dans une partie où elle n'est
# pas annoncée fait échouer — même annoncée pour la partie voisine, sans quoi les axes
# de Bravo admis dans l'onglet `fiche` passeraient aussi dans l'onglet `activite` — ; une
# sentinelle annoncée qui ne passe plus aussi.
#
# Bâti par BALAYAGE le 2026-10-07, pas de mémoire. Il décrit l'existant, y compris ce qui
# attend une décision : c'est un instrument d'inventaire, pas un arbitrage. Il ne reste
# ici que des LIMITES écrites : le seul écart « à trancher » qu'il ait porté — l'export
# d'album, dont le vocabulaire variait selon qui clique — a été tranché le jour même et
# retiré, le cliquet ayant crié sur la déclaration devenue périmée.
Ecart = namedtuple("Ecart", "parties raison")

_AXES_BRAVO = frozenset({"dom-bravo-9102", "dim-bravo-9202", "val-bravo-9302",
                         "valg-sous-bravo-9312", "valp-bravo-9311"})

_LIMITE_FICHE = (
    "LIMITE ÉCRITE, non tranchée (AUTH-11, 2026-09-23 ; `docs/export-metadonnees.md`, "
    "§ Portée d'une collection). Le bloc `vocabulaire` de la fiche de description nomme "
    "toutes les dimensions, toutes leurs valeurs et tous les domaines de l'INSTANCE : "
    "ses chiffres décrivent une couverture, et dire s'ils portent sur l'instance ou sur "
    "le périmètre est une question de modèle, pas une clause à poser. Les TAGS n'y "
    "passent pas — ils y sont comptés, jamais nommés — et c'est pourquoi la sentinelle "
    "de tag n'est pas dans cette liste : la voir apparaître serait une fuite neuve.")

_LIMITE_PORTEE = (
    "LIMITE ÉCRITE (AUTH-11, 2026-09-24, case « taire les charges ») : `activite.params`, "
    "`portee` et `comptes` sortent au grain CORPUS. Aucun contenu de travail, mais "
    "`portee` désigne des albums et des planches qu'on ne dépose pas. La décision "
    "« un acte n'est pas re-scopable » n'a pas été rouverte pour les runs.")

_PROV = _LIMITE_PORTEE + " La sérialisation PROV-O publie les mêmes trois colonnes."

ECARTS = {
    ("route", _DESC, "json"): Ecart({"": _AXES_BRAVO}, _LIMITE_FICHE),
    ("outil", "description_collection.py", "--json"): Ecart({"": _AXES_BRAVO}, _LIMITE_FICHE),
    # Le classeur embarque la fiche : la même limite, par ricochet, dans UN onglet — et
    # la portée des runs dans un AUTRE. Chacune n'est admise que chez elle.
    ("route", _META, "xlsx"): Ecart(
        {"fiche": _AXES_BRAVO, "activite": {PORTEE_AILLEURS}},
        _LIMITE_FICHE + " — onglet `fiche`. Et, onglet `activite` : " + _LIMITE_PORTEE),
    ("outil", "metadonnees_collection.py", "--xlsx"): Ecart(
        {"m.xlsx/fiche": _AXES_BRAVO, "m.xlsx/activite": {PORTEE_AILLEURS}},
        _LIMITE_FICHE + " — onglet `fiche`. Et, onglet `activite` : " + _LIMITE_PORTEE),
    ("route", _META, "zip"): Ecart({"activite.csv": {PORTEE_AILLEURS}}, _LIMITE_PORTEE),
    ("outil", "metadonnees_collection.py", "--zip"): Ecart(
        {"m.zip/activite.csv": {PORTEE_AILLEURS}}, _LIMITE_PORTEE),
    ("outil", "metadonnees_collection.py", "--csv-dir"): Ecart(
        {"activite.csv": {PORTEE_AILLEURS}}, _LIMITE_PORTEE),
    ("outil", "provenance_export.py", "--out-dir"): Ecart(
        {"provenance.json": {PORTEE_AILLEURS}}, _PROV),
    ("outil", "provenance_export.py", "(stdout)"): Ecart({"": {PORTEE_AILLEURS}}, _PROV),
}


# --------------------------------------------------------------------------- #
# Les routes HORS du contrôle, chacune avec sa raison
# --------------------------------------------------------------------------- #
_TRANSVERSAL = (
    "Export TRANSVERSAL : il ne sort au titre d'aucune collection, il traverse toutes "
    "celles que la PERSONNE a le droit d'exporter (`Portee.pour_export`, DROIT-2). « Le "
    "vocabulaire de A et de A seule » n'a pas d'objet sans A. Sa règle est celle de la "
    "personne, et `test_droit_export` la joue (`test_la_case_sur_a_n_emporte_pas_b`, "
    "`test_ce_qui_sort_suit_la_portee_d_export_pas_celle_de_lecture`).")
_BASE_ENTIERE = (
    "La base ENTIÈRE, par construction (DROIT-1) : toutes les collections, tous les "
    "termes, le journal avec ses charges. Une sauvegarde partielle ne restaure pas une "
    "instance. Réservée aux administrateurs ; ce n'est pas la sortie d'UNE collection.")

HORS_CONTROLE = {
    ("GET", "/api/recherche/export.csv"): _TRANSVERSAL,
    ("GET", "/api/analyse/frequences.csv"): _TRANSVERSAL,
    ("GET", "/api/analyse/concordance.csv"): _TRANSVERSAL,
    ("GET", "/api/analyse/comparaison.csv"): _TRANSVERSAL,
    ("GET", "/api/analyse/croisement.csv"): _TRANSVERSAL,
    ("GET", "/api/analyse/accord.csv"): _TRANSVERSAL
        + " Celui-ci ne porte d'ailleurs que des tokens et des taux, aucun terme.",
    ("GET", "/api/analyse/accord-inter.csv"): _TRANSVERSAL
        + " Celui-ci nomme des PERSONNES et non des termes : c'est AUTH-5 qui le garde.",
    ("GET", "/api/sauvegarde"): _BASE_ENTIERE,
    ("POST", "/api/sharedocs/deposer-sauvegarde"): _BASE_ENTIERE
        + " Même artefact que `GET /api/sauvegarde`, envoyé sur ShareDocs.",
}

# Les routes jouées qui ne sont PAS sous `/depot/` : comment les appeler, et les
# paramètres de requête qu'on leur connaît. Un paramètre neuf fait échouer — il peut
# être un format, ou un interrupteur qui fait sortir davantage.
JOUEES_A_PART = {
    ("GET", "/api/export/json"): {"album_id", "collection_id"},
    ("GET", "/api/export/csv"): {"album_id", "collection_id"},
    ("GET", "/api/export/tei"): {"album_id", "collection_id"},
    ("POST", "/api/figures"): set(),
}

# Ce que le cliquet sait faire varier sur une route de dépôt. `format` est énuméré
# depuis le contrat de la route ; les deux autres sont joués dans chaque position.
PARAMETRES_DEPOT = {
    "format": None,
    "verbatim": (False, True),
    "base_url": ("", ADRESSE_IMAGES),      # sans adresse, le manifeste sort en APERÇU
}
# Le corps du dépôt ShareDocs : les mêmes réglages, plus la destination.
CHAMPS_DEPOSER = {"quoi", "format", "verbatim", "base_url", "dossier", "compte"}


# --------------------------------------------------------------------------- #
# Les outils : joués, ou déclarés hors du contrôle
# --------------------------------------------------------------------------- #
# Une entrée par ENTRÉE que l'`argparse` de l'outil déclare, relevée par AST
# (`_entrees_argparse`) : option longue ou courte, argument positionnel, sous-commande —
# préfixée de la sous-commande qui la porte (« montrer id »), ou de la fonction d'aide
# qui la pose (« (_ajouter_descripteurs) --licence »). « sortie » : elle produit un
# artefact, joué et jugé. « variante » : un interrupteur de contenu, joué dans ses deux
# positions. Tout autre mot est une raison de ne pas la jouer. Une entrée absente d'ici
# fait échouer : c'est par une option ajoutée — `--xlsx`, pour AUTH-1 — qu'une sortie
# neuve entre sans être regardée.
SORTIE, VARIANTE = "sortie", "variante"
SYS_ARGV = "(sys.argv)"                 # l'outil lit sa ligne de commande à la main
NON_LITTERAL = "<non littéral>"         # un nom d'entrée calculé : illisible par AST
_ECRIT = "appartient à une sous-commande qui ÉCRIT en base : rien n'en sort"
_CIBLE = "l'identifiant de la collection visée : une cible, pas une sortie"
OUTILS_JOUES = {
    "description_collection.py": {
        "--json": SORTIE, "--csv": SORTIE,
        "--collection": "le périmètre : toujours posé, c'est l'objet du contrôle"},
    "metadonnees_collection.py": {
        "--json": SORTIE, "--csv-dir": SORTIE, "--zip": SORTIE, "--xlsx": SORTIE,
        "--verbatim": VARIANTE,
        "--collection": "le périmètre : toujours posé, c'est l'objet du contrôle"},
    "iiif_manifest.py": {
        "--out-dir": SORTIE, "--verbatim": VARIANTE,
        "--base-url": "préfixe d'URI des identifiants : ne change pas ce qui est publié",
        "--collection": "le périmètre : toujours posé, c'est l'objet du contrôle"},
    "crosswalk_depot.py": {
        "--out-dir": SORTIE,
        "--publisher": "une mention de la notice, fournie par l'appelant",
        "--annee-depot": "une mention de la notice, fournie par l'appelant",
        "--collection": "le périmètre : toujours posé, c'est l'objet du contrôle"},
    "provenance_export.py": {"--out-dir": SORTIE},
    "gerer_collections.py": {
        "lister": SORTIE, "montrer": SORTIE, "montrer id": _CIBLE,
        # Chaque entrée est rangée sous SA sous-commande : `--albums` est classée pour
        # `creer`, `ajouter` et `retirer`, et la poser sur `montrer` ferait une entrée
        # neuve (« montrer --albums »), donc un échec. Aplaties, les options d'écriture
        # auraient couvert la même option posée sur une sous-commande de lecture.
        **{e: _ECRIT for e in (
            "creer", "creer --nom", "creer --albums", "creer --proprietaire",
            "creer --proprietaire-groupe", "creer +_ajouter_descripteurs()",
            "modifier", "modifier id", "modifier --nom",
            "modifier +_ajouter_descripteurs()",
            "ajouter", "ajouter id", "ajouter --albums",
            "retirer", "retirer id", "retirer --albums",
            "supprimer", "supprimer id",
            "(_ajouter_descripteurs) --description", "(_ajouter_descripteurs) --licence",
            "(_ajouter_descripteurs) --statut", "(_ajouter_descripteurs) --date-embargo",
            "(_ajouter_descripteurs) --base-legale", "(_ajouter_descripteurs) --responsable",
            "(_ajouter_descripteurs) --date-debut", "(_ajouter_descripteurs) --date-fin")}},
}
# Les outils qui écrivent sur la sortie standard quand aucune option de sortie n'est
# donnée. Ce n'est pas une option : `argparse` ne la déclare pas, donc l'AST ne la voit
# pas — et c'est une sortie quand même (le défaut de `iiif_manifest` est un manifeste).
SORTIE_PAR_DEFAUT = {"iiif_manifest.py", "crosswalk_depot.py", "provenance_export.py"}

OUTILS_HORS = {
    "_commun.py": "Bibliothèque partagée, sans `main()` : elle ne produit rien.",
    "importer_vocabulaire.py":
        "ENTRE du vocabulaire, n'en sort pas. Son `--collection` désigne la portée où "
        "les termes naissent ; son compte rendu ne nomme que les lignes du tableur "
        "qu'on lui donne. Ses refus sont gardés par `test_import_vocabulaire`.",
    "dictionnaire_xlsx.py":
        "Met en classeur le DICTIONNAIRE des métadonnées, c'est-à-dire un fichier "
        "markdown du dépôt (`docs/dictionnaire-metadonnees.md`). N'ouvre pas la base.",
    "rapport_accord.py":
        "Rapport modèle↔humain au grain CORPUS, sans périmètre de collection : des "
        "tokens, des catégories grammaticales et des taux. Il ne lit ni les tags, ni "
        "les attributs, ni le journal.",
    "rapport_accord_inter.py":
        "Rapport inter-annotateurs au grain CORPUS, sans périmètre de collection. Il "
        "lit le journal, mais les seuls actes `token_correction` (lemme, POS, morpho) "
        "— jamais une charge d'annotation. Il nomme des PERSONNES : c'est AUTH-5 qui "
        "le garde, pas ce cliquet.",
    "plafonds_dependances.py": "Croise les verrous du dépôt et les paquets installés. "
                               "N'ouvre pas la base.",
    "identite_pile.py": "Versions, empreintes des verrous, commit servi. N'ouvre pas "
                        "la base.",
    "verifier_moteurs.py": "Vérifie la présence des moteurs ML. N'ouvre pas la base.",
    "pdf_check.py": "Contrôle un PDF fourni en argument. N'ouvre pas la base.",
    "sharedocs_check.py": "Vérifie une connexion WebDAV. N'ouvre pas la base.",
    "valider_iiif.py": "Valide des manifestes FOURNIS en argument. N'ouvre pas la base.",
    "mesurer_reflow.py": "Mesure des rectangles dans un navigateur. N'ouvre pas la base.",
    "faux_proxy_auth.py": "Faux proxy de développement : relaie du HTTP. N'ouvre pas "
                          "la base.",
    "reindex_nlp.py": "Maintenance : régénère tokens et index. Compte rendu = décompte.",
    "reindex_materiel.py": "Maintenance : relit les masters. Compte rendu = décompte.",
    "regenerer_derives.py": "Maintenance : réécrit les dérivés. Compte rendu = décompte.",
    "normaliser_casse.py":
        "Maintenance (NLP-3) : son aperçu recopie des RÉPLIQUES, du texte de bulle — "
        "jamais un terme de vocabulaire ni une charge du journal.",
    "semer_demo.py": "SÈME un corpus de démonstration par l'API : il écrit, il ne "
                     "sort rien.",
    "regenerer_exemples.py":
        "Écrit `docs/exemples/` depuis un corpus JETABLE qu'il sème lui-même, jamais "
        "depuis la base réelle. Le jouer réécrirait le dépôt versionné pendant la suite.",
}


# Les entrées `argparse` de chaque outil HORS du contrôle, FIGÉES. Sa raison d'être hors
# (« n'exporte rien », « rapport au grain corpus ») a été écrite pour CES entrées : une
# option neuve — un `--dump-vocabulaire` sur un rapport — la rend fausse sans la toucher.
# Elle fait donc échouer, et celui qui l'ajoute relit la raison avant de l'inscrire ici.
ENTREES_HORS = {
    "_commun.py": set(),
    "importer_vocabulaire.py": {"fichier", "--collection", "--dry-run"},
    "dictionnaire_xlsx.py": {"--out", "--source"},
    "rapport_accord.py": {"--json", "--csv", "--modele"},
    "rapport_accord_inter.py": {"--json", "--csv"},
    "plafonds_dependances.py": {"--tous"},
    "identite_pile.py": {"--identite", "--json"},
    "verifier_moteurs.py": {"--exiger", "--json"},
    "pdf_check.py": {SYS_ARGV},
    "sharedocs_check.py": {SYS_ARGV},
    "valider_iiif.py": {"chemins"},
    "mesurer_reflow.py": {SYS_ARGV},
    "faux_proxy_auth.py": set(),
    "reindex_nlp.py": set(),
    "reindex_materiel.py": {"--dry-run", "--force"},
    "regenerer_derives.py": {"--album", "--dry-run", "--mode", "--planche", "--toutes"},
    "normaliser_casse.py": {"--album", "--dry-run", "--planche"},
    "semer_demo.py": set(),
    "regenerer_exemples.py": set(),
}


def _outils() -> list:
    return sorted(p.name for p in TOOLS.glob("*.py"))


class _Releveur(ast.NodeVisitor):
    """Relève les entrées `argparse` d'un source, chacune sous la sous-commande qui la porte.

    Suit, fonction par fonction et dans l'ordre du source, la variable qui reçoit un
    `add_parser("x")` : un `add_argument` appelé sur elle appartient à « x ». Un
    `add_argument` appelé sur un PARAMÈTRE de fonction appartient à cette fonction d'aide,
    et l'appel de l'aide sur un sous-analyseur est relevé à son tour (« x +aide() ») —
    sans quoi brancher l'aide sur une autre sous-commande ne se verrait pas.
    """

    def __init__(self):
        self.vues = set()
        self.fonctions = []            # pile : (nom, paramètres)
        self.portee = {}               # (fonction, variable) -> sous-commande

    def _fn(self):
        return self.fonctions[-1][0] if self.fonctions else ""

    def visit_FunctionDef(self, n):
        self.fonctions.append((n.name, {a.arg for a in n.args.args}))
        self.generic_visit(n)
        self.fonctions.pop()

    visit_AsyncFunctionDef = visit_FunctionDef

    @staticmethod
    def _sous_commande(n):
        for x in ast.walk(n):
            if (isinstance(x, ast.Call) and isinstance(x.func, ast.Attribute)
                    and x.func.attr == "add_parser" and x.args
                    and isinstance(x.args[0], ast.Constant)):
                return x.args[0].value
        return None

    def visit_Assign(self, n):
        self.visit(n.value)
        sc = self._sous_commande(n.value)
        for c in n.targets:
            if isinstance(c, ast.Name):
                if sc is not None:
                    self.portee[(self._fn(), c.id)] = sc
                else:
                    self.portee.pop((self._fn(), c.id), None)

    def _ou(self, recepteur):
        if not isinstance(recepteur, ast.Name):
            return ""
        if self.fonctions and recepteur.id in self.fonctions[-1][1]:
            return f"({self._fn()})"
        return self.portee.get((self._fn(), recepteur.id), "")

    @staticmethod
    def _litteral(a):
        return a.value if isinstance(a, ast.Constant) and isinstance(a.value, str) \
            else NON_LITTERAL

    def visit_Call(self, n):
        if isinstance(n.func, ast.Attribute) and n.func.attr == "add_parser":
            self.vues.add(self._litteral(n.args[0]) if n.args else NON_LITTERAL)
        elif isinstance(n.func, ast.Attribute) and n.func.attr == "add_argument":
            ou = self._ou(n.func.value)
            for a in n.args or [None]:
                self.vues.add(f"{ou} {self._litteral(a)}".strip())
        elif isinstance(n.func, ast.Name):
            for a in n.args:
                if isinstance(a, ast.Name) and (self._fn(), a.id) in self.portee:
                    self.vues.add(f"{self.portee[(self._fn(), a.id)]} +{n.func.id}()")
        self.generic_visit(n)


def _entrees_argparse(outil: str) -> set:
    """Les ENTRÉES que l'outil déclare, lues dans son SOURCE : options longues et courtes,
    arguments positionnels, sous-commandes — chacune sous la sous-commande qui la porte.

    Par AST et non en important l'outil : la moitié d'entre eux ouvrent la base ou
    reconfigurent la sortie standard à l'import de leur `main`.

    **Ce que ce relevé ne tient PAS, écrit plutôt que promis.**
    - Un outil qui lit `sys.argv` À LA MAIN (`pdf_check`, `sharedocs_check`,
      `mesurer_reflow`) : seul le FAIT qu'il le lise est figé (`SYS_ARGV`). Un drapeau de
      plus dans un tel outil ne se voit pas ; un outil `argparse` qui se mettrait à lire
      `sys.argv` en plus, si.
    - Un nom d'entrée CALCULÉ (`add_argument(*noms)`) : il sort en `NON_LITTERAL`, qui
      n'est classé nulle part, donc il fait échouer — c'est un refus, pas une lecture.
    - Une sortie qui ne dépend d'AUCUNE entrée : un fichier écrit d'office par un outil
      hors du contrôle ne change aucune entrée, et rien ne le voit ici.
    - `parse_known_args`, une variable d'environnement, un fichier de configuration.
    """
    source = (TOOLS / outil).read_text(encoding="utf-8")
    r = _Releveur()
    r.visit(ast.parse(source))
    return r.vues | ({SYS_ARGV} if "sys.argv" in source else set())


def _sorties_d_outil(outil: str) -> list:
    """Les sorties à jouer pour un outil : ses options « sortie », plus son défaut."""
    s = sorted(e for e, role in OUTILS_JOUES[outil].items() if role == SORTIE)
    return s + (["(stdout)"] if outil in SORTIE_PAR_DEFAUT else [])


# --------------------------------------------------------------------------- #
# L'énumération des routes
# --------------------------------------------------------------------------- #
def _est_une_sortie(route) -> bool:
    """Une route qui produit un fichier, ou qui ENVOIE hors de l'instance.

    La première moitié est le critère de DROIT-2, importé et non recopié : deux
    définitions de « porte de sortie » finiraient par diverger. La seconde le complète —
    le dépôt d'une sauvegarde sur ShareDocs ne rend aucun fichier à l'appelant, donc ce
    critère-là ne le voit pas, et c'est pourtant la base entière qui part.
    """
    return _sort_un_fichier(route) or "sharedocs.upload(" in inspect.getsource(route.endpoint)


def _routes_de_sortie() -> dict:
    routes = inventaire_routes.routes_api()
    inventaire_routes.exiger_plancher(len(routes), "le cliquet des sorties d'une "
                                                   "collection (AUTH-11)")
    return {(m, r.path): r for r in routes if _est_une_sortie(r) for m in r.methods}


def _parametres(route) -> dict:
    """Les paramètres de REQUÊTE d'une route, par nom."""
    return {p.name: p for p in route.dependant.query_params}


def _formats_declares(route) -> tuple:
    """Les formats qu'une route de dépôt sait rendre, lus dans SON contrat.

    Le `pattern` du paramètre `format` est ce que la route accepte réellement (et ce que
    `/docs` publie) ; une route sans ce paramètre ne rend que ce que `_FORMATS` annonce
    pour elle. Les deux sources sont confrontées par le test d'énumération : jouer un
    format que la route refuse, ou en ignorer un qu'elle accepte, serait la même faute.
    """
    import routes.depot as depot
    quoi = route.path.rsplit("/", 1)[1]
    p = _parametres(route).get("format")
    if p is None:
        return tuple(depot._FORMATS.get(quoi, ()))
    motif = next((getattr(m, "pattern", None) for m in p.field_info.metadata
                  if getattr(m, "pattern", None)), None)
    trouve = re.fullmatch(r"\^\(([a-z0-9|]+)\)\$", motif or "")
    assert trouve, (
        f"{route.path} : le paramètre `format` n'a pas la forme `^(a|b|c)$` "
        f"({motif!r}). Le cliquet ne sait plus énumérer ses formats — donc il n'en "
        "jouerait aucun, et passerait au vert.")
    return tuple(trouve.group(1).split("|"))


def _routes_de_depot(sorties: dict) -> dict:
    return {cle: r for cle, r in sorties.items() if DEPOT in cle[1]}


# --------------------------------------------------------------------------- #
# Le cliquet d'ÉNUMÉRATION — sans base, sans appel : la forme du contrat
# --------------------------------------------------------------------------- #
def test_toute_sortie_d_une_collection_est_jouee_ou_declaree():
    """Une sortie nouvelle, ni jouée ni déclarée, fait ÉCHOUER — et une déclaration qui
    ne désigne plus rien aussi."""
    import routes.depot as depot
    sorties = _routes_de_sortie()
    depots = _routes_de_depot(sorties)
    fautes = []
    assert depots, "aucune route de dépôt trouvée : le cliquet ne mesurerait rien"

    # (1) Toute route de sortie est tranchée. Une route `/depot/` en GET est jouée
    #     d'office ; en POST, seule la voie ShareDocs sait l'être.
    jouees = set(JOUEES_A_PART) | {c for c in depots if c[0] == "GET"} | {("POST", DEPOSER)}
    for cle in sorted(sorties):
        if cle not in jouees and cle not in HORS_CONTROLE:
            fautes.append(f"{cle[0]} {cle[1]} fait sortir quelque chose et n'est ni "
                          "jouée ni déclarée hors du contrôle")
        if cle in jouees and cle in HORS_CONTROLE:
            fautes.append(f"{cle[0]} {cle[1]} est à la fois jouée et déclarée hors")
    for cle, raison in sorted(HORS_CONTROLE.items()):
        assert raison and len(raison) > 40, f"{cle} est hors du contrôle sans raison écrite"
        if cle not in sorties:
            fautes.append(f"{cle[0]} {cle[1]} est déclarée hors du contrôle mais n'est "
                          "plus une sortie : retirer la déclaration")
    for cle in sorted(set(JOUEES_A_PART) | {("POST", DEPOSER)}):
        if cle not in sorties:
            fautes.append(f"{cle[0]} {cle[1]} est dite jouée mais n'est plus une sortie")

    # (2) Chaque route de dépôt, dans chaque format de SON contrat, a son attendu ; et
    #     aucun de ses paramètres n'est inconnu du cliquet.
    cles_route = set()
    for (methode, chemin), r in sorted(depots.items()):
        if methode != "GET":
            continue
        quoi = chemin.rsplit("/", 1)[1]
        formats = _formats_declares(r)
        if not formats:
            fautes.append(f"{chemin} : aucun format connu (ni paramètre `format`, ni "
                          f"entrée « {quoi} » dans `routes.depot._FORMATS`)")
        if tuple(depot._FORMATS.get(quoi, ())) != formats:
            fautes.append(f"{chemin} accepte {formats} et `_FORMATS[{quoi!r}]` annonce "
                          f"{depot._FORMATS.get(quoi)} : la voie ShareDocs et le "
                          "téléchargement ne rendraient plus les mêmes formats")
        for nom in sorted(set(_parametres(r)) - set(PARAMETRES_DEPOT)):
            fautes.append(f"{chemin} : paramètre `{nom}` inconnu du cliquet — il peut "
                          "être un format ou faire sortir davantage. L'ajouter à "
                          "PARAMETRES_DEPOT avec les positions à jouer.")
        for f in formats:
            cles_route.add(("route", chemin, f))
    quois = {c[1].rsplit("/", 1)[1] for c in depots if c[0] == "GET"}
    if set(depot._FORMATS) != quois:
        fautes.append(f"`_FORMATS` connaît {sorted(depot._FORMATS)} et les routes GET de "
                      f"dépôt {sorted(quois)} : un export que l'une des deux voies — le "
                      "téléchargement, ShareDocs — est seule à connaître ne serait pas "
                      "joué par l'autre")
    champs = set(main.DeposerExportIn.model_fields)
    if champs != CHAMPS_DEPOSER:
        fautes.append(f"le corps du dépôt ShareDocs a changé ({sorted(champs ^ CHAMPS_DEPOSER)}) "
                      ": dire comment jouer le champ neuf, ou le retirer de CHAMPS_DEPOSER")
    for cle, r in sorted(sorties.items()):
        if cle in JOUEES_A_PART:
            ecart = set(_parametres(r)) ^ JOUEES_A_PART[cle]
            if ecart:
                fautes.append(f"{cle[0]} {cle[1]} : paramètres de requête {sorted(ecart)} "
                              "en plus ou en moins de ce que le cliquet sait jouer")

    # (3) Tout outil est joué ou déclaré, et toute option d'un outil joué est classée.
    reels = set(_outils())
    for o in sorted(reels - set(OUTILS_JOUES) - set(OUTILS_HORS)):
        fautes.append(f"tools/{o} n'est ni joué ni déclaré hors du contrôle")
    for o in sorted((set(OUTILS_JOUES) | set(OUTILS_HORS)) - reels):
        fautes.append(f"tools/{o} est déclaré mais n'existe plus")
    for o in sorted(set(OUTILS_JOUES) & set(OUTILS_HORS)):
        fautes.append(f"tools/{o} est à la fois joué et déclaré hors")
    for o, raison in sorted(OUTILS_HORS.items()):
        assert raison and len(raison) > 40, f"tools/{o} est hors du contrôle sans raison"
        if o in reels and "--collection" in _entrees_argparse(o) \
                and "--collection" not in raison:
            fautes.append(f"tools/{o} prend `--collection` et sa raison d'être hors du "
                          "contrôle n'en dit rien : un outil qui sait borner une "
                          "collection est jouable, ou sa raison l'explique")
        if o in reels:
            vues, figees = _entrees_argparse(o), ENTREES_HORS.get(o)
            if figees is None:
                fautes.append(f"tools/{o} est hors du contrôle sans que ses entrées "
                              "soient figées (ENTREES_HORS)")
            elif vues != figees:
                fautes.append(
                    f"tools/{o} est déclaré hors du contrôle, et ses entrées ont changé "
                    f"— en plus {sorted(vues - figees)}, en moins {sorted(figees - vues)}. "
                    "Relire sa raison dans OUTILS_HORS : elle a été écrite pour les "
                    "entrées d'avant. Puis mettre ENTREES_HORS à jour, ou jouer l'outil.")
    for o in sorted(set(ENTREES_HORS) - set(OUTILS_HORS)):
        fautes.append(f"tools/{o} a des entrées figées sans être déclaré hors du contrôle")
    cles_outil = set()
    for o, classees in sorted(OUTILS_JOUES.items()):
        if o not in reels:
            continue
        vues = _entrees_argparse(o)
        for e in sorted(vues - set(classees)):
            fautes.append(f"tools/{o} : `{e}` n'est pas classée — sortie, variante, ou "
                          "une raison de ne pas la jouer")
        for e in sorted(set(classees) - vues):
            fautes.append(f"tools/{o} : `{e}` est classée mais l'outil ne la déclare plus")
        cles_outil |= {("outil", o, s) for s in _sorties_d_outil(o)}
    for o in sorted(SORTIE_PAR_DEFAUT - set(OUTILS_JOUES)):
        fautes.append(f"tools/{o} figure dans SORTIE_PAR_DEFAUT sans être joué")

    # (4) Les attendus et les écarts désignent des sorties RÉELLES, et chacune a le sien.
    a_part = {k for k in ATTENDUS if k[0] == "route"
              and any(k[1] == c[1] for c in JOUEES_A_PART)}
    for c in sorted(JOUEES_A_PART):
        if not any(k[1] == c[1] for k in a_part):
            fautes.append(f"{c[0]} {c[1]} est dite jouée et n'a aucun attendu")
    reelles = cles_route | cles_outil | a_part
    for cle in sorted(reelles - set(ATTENDUS)):
        fautes.append(f"{cle} est une sortie réelle sans attendu : dire ce qu'elle "
                      "doit porter (ATTENDUS), le cliquet la jouera")
    for cle in sorted(set(ATTENDUS) - reelles):
        fautes.append(f"{cle} a un attendu mais n'est plus une sortie")
    for cle, e in sorted(ECARTS.items()):
        assert e.raison and len(e.raison) > 40, f"{cle} : écart sans raison écrite"
        assert e.parties and all(e.parties.values()), f"{cle} : écart qui ne désigne rien"
        if cle not in ATTENDUS:
            fautes.append(f"{cle} déclare un écart et n'est pas une sortie jouée")
    for cle, a in sorted(ATTENDUS.items()):
        assert a.raison, f"{cle} : attendu sans raison"
        assert a.porte or a.temoin, (
            f"{cle} n'annonce ni vocabulaire ni témoin : rien ne prouverait qu'on a LU "
            "l'artefact, et un lecteur vacant le dirait propre")

    assert not fautes, ("Le cliquet des sorties d'une collection ne décrit plus la "
                        "réalité.\n  " + "\n  ".join(fautes))


# --------------------------------------------------------------------------- #
# Le balayage
# --------------------------------------------------------------------------- #
def _variantes(route) -> list:
    """Chaque combinaison des paramètres d'une route de dépôt, `format` mis à part."""
    combis = [{}]
    for nom in sorted(set(_parametres(route)) - {"format"}):
        positions = PARAMETRES_DEPOT.get(nom)
        assert positions, f"{route.path} : `{nom}` — le cliquet ne sait pas le jouer"
        combis = [{**c, nom: p} for c in combis for p in positions]
    return combis


def _balayer_routes(client, seme, collection_id: int, monkeypatch=None, *,
                    album: bool = True) -> list:
    """Joue chaque route de sortie d'UNE collection. Rend [(clé, variante, parties)].

    `album=False` laisse de côté l'export d'album et la figure : l'album et la région du
    décor sont ceux d'Alpha, et les demander au titre de Bravo répond 404.

    Avec `monkeypatch`, joue aussi la voie ShareDocs, l'envoi WebDAV étant remplacé par
    une capture : c'est ce qui SERAIT parti qu'on juge, contre l'attendu du téléchargement
    du même export — les deux voies promettent le même octet.
    """
    sorties = _routes_de_sortie()
    vus = []

    def appeler(cle, variante, methode, url, **kw):
        rep = client.request(methode, url, headers=ADMIN, **kw)
        if cle[2] == "xlsx" and rep.status_code == 503 and not _xlsx_lisible():
            return                                  # extra absent : dit par le dernier test
        assert rep.status_code == 200, (cle, variante, rep.status_code, rep.text[:300])
        vus.append((cle, variante, _parties(rep.content)))

    for (methode, chemin), r in sorted(_routes_de_depot(sorties).items()):
        if methode != "GET":
            continue
        url = chemin.replace("{collection_id}", str(collection_id))
        for f in _formats_declares(r):
            for v in _variantes(r):
                params = {**v, **({"format": f} if "format" in _parametres(r) else {})}
                appeler(("route", chemin, f), json.dumps(v, sort_keys=True), "GET", url,
                        params=params)

    for fmt in ("json", "csv", "tei") if album else ():
        appeler(("route", f"/api/export/{fmt}", fmt), "au titre de la collection", "GET",
                f"/api/export/{fmt}",
                params={"album_id": seme["album_id"], "collection_id": collection_id})
    if album:
        import figure
        appeler(("route", "/api/figures", "zip"), "toutes les mentions", "POST",
                "/api/figures", json={"regions": [seme["region_id"]],
                                      "champs": list(figure.CHAMPS),
                                      "collection_id": collection_id})

    if monkeypatch is not None:
        import pipeline.sharedocs as sd
        import routes.depot as depot
        capte = {}
        monkeypatch.setattr(
            sd, "upload",
            lambda chemin, data, *, principal, compte=None:
                capte.update(data=data) or {"chemin": chemin, "compte": "instance",
                                            "user": "u"})
        url = DEPOSER.replace("{collection_id}", str(collection_id))
        for quoi, formats in sorted(depot._FORMATS.items()):
            for f in formats:
                reglages = [{"verbatim": vb, "base_url": bu}
                            for vb in PARAMETRES_DEPOT["verbatim"]
                            for bu in PARAMETRES_DEPOT["base_url"]]
                for reglage in reglages:
                    capte.clear()
                    rep = client.post(url, headers=ADMIN, json={
                        "quoi": quoi, "format": f, "dossier": "D", **reglage})
                    if f == "xlsx" and rep.status_code == 503 and not _xlsx_lisible():
                        continue
                    assert rep.status_code == 200, (quoi, f, rep.status_code, rep.text[:300])
                    assert capte.get("data"), f"dépôt {quoi}/{f} : rien n'est parti"
                    # Jugé contre l'attendu du TÉLÉCHARGEMENT du même export.
                    cle = ("route", DEPOSER.replace("deposer", quoi), f)
                    vus.append((cle, "voie ShareDocs, " + json.dumps(reglage, sort_keys=True),
                                _parties(capte["data"])))
    return vus


def _invocations(seme, t: Path) -> list:
    """(clé, variante, outil, arguments, chemin produit à relire ou None = stdout)."""
    a = ["--collection", str(seme["alpha"]["id"])]
    inv = []

    def ajouter(outil, sortie, args, produit=None, variantes=((),)):
        for extra in variantes:
            inv.append((("outil", outil, sortie), " ".join(extra) or "—", outil,
                        [*args, *extra], produit))

    verb = ((), ("--verbatim",))
    ajouter("description_collection.py", "--json", ["--json", "-", *a])
    ajouter("description_collection.py", "--csv", ["--csv", "-", *a])
    ajouter("metadonnees_collection.py", "--json", ["--json", "-", *a], variantes=verb)
    ajouter("metadonnees_collection.py", "--csv-dir", ["--csv-dir", str(t / "csv"), *a],
            t / "csv", verb)
    ajouter("metadonnees_collection.py", "--zip", ["--zip", str(t / "m.zip"), *a],
            t / "m.zip", verb)
    ajouter("metadonnees_collection.py", "--xlsx", ["--xlsx", str(t / "m.xlsx"), *a],
            t / "m.xlsx", verb)
    ajouter("iiif_manifest.py", "--out-dir",
            ["--base-url", ADRESSE_IMAGES, "--out-dir", str(t / "iiif"), *a],
            t / "iiif", verb)
    ajouter("iiif_manifest.py", "(stdout)", ["--base-url", ADRESSE_IMAGES, *a],
            variantes=verb)
    ajouter("crosswalk_depot.py", "--out-dir", ["--out-dir", str(t / "cw"), *a], t / "cw")
    ajouter("crosswalk_depot.py", "(stdout)", a)
    # Pas de `--collection` : le journal sort au grain corpus, et c'est écrit. On le joue
    # quand même — c'est lui qui publiait les charges.
    ajouter("provenance_export.py", "--out-dir", ["--out-dir", str(t / "prov")], t / "prov")
    ajouter("provenance_export.py", "(stdout)", [])
    ajouter("gerer_collections.py", "lister", ["lister"])
    ajouter("gerer_collections.py", "montrer", ["montrer", str(seme["alpha"]["id"])])
    return inv


def _balayer_outils(seme, tmp_path) -> list:
    """Lance chaque outil en SOUS-PROCESSUS, comme un humain au shell, et relit ce qu'il
    a écrit — sortie standard, sortie d'erreur, et fichiers."""
    vus = []
    env = {**os.environ, "BD_DB_PATH": str(seme["db"]), "BD_DATA_DIR": str(seme["data"])}
    for cle, variante, outil, args, produit in _invocations(seme, tmp_path):
        if cle[2] == "--xlsx" and not _xlsx_lisible():
            continue                                # extra absent : dit par le dernier test
        p = subprocess.run([sys.executable, str(TOOLS / outil), *args],
                           cwd=str(REPO_ROOT), env=env, capture_output=True)
        assert p.returncode == 0, (cle, variante, p.stderr.decode("utf-8", "replace")[-900:])
        # Les deux flux sont du TEXTE de console, pas un conteneur : aucun membre binaire
        # ne peut s'y cacher, et la sortie d'erreur d'un outil qui n'a pas forcé l'UTF-8
        # sort en cp1252 sous Windows. D'où `replace` ici, et ici seulement — les
        # sentinelles sont en ASCII, un accent mal décodé ne les masque pas.
        parties = {"": p.stdout.decode("utf-8", "replace"),
                   "(stderr)": p.stderr.decode("utf-8", "replace")}
        if produit is not None:
            parties.update(_fichiers(Path(produit)))
        vus.append((cle, variante, parties))
    return vus


def _presents(mots, texte: str) -> set:
    """Recherche insensible à la casse : une clé de tri, un slug ou une normalisation
    minusculent couramment, et la sentinelle y échapperait sans rien dire (AUTH-5)."""
    bas = texte.lower()
    return {m for m in mots if m.lower() in bas}


def _juger(vus: list, propre: list, interdits: list) -> list:
    """Chaque artefact joué : propre, ou son écart déclaré au plus près — et rien de plus."""
    fautes = []
    for cle, variante, parties in vus:
        ou = f"{cle[1]} [{cle[2]}] ({variante})"
        attendu = ATTENDUS.get(cle)
        if attendu is None:
            # Fouillée quand même : la faute dit d'emblée ce que la sortie neuve emporte.
            fuite = _presents(interdits, "\n".join(parties.values()))
            fautes.append(f"{ou} : jouée sans attendu déclaré"
                          + (f" — et elle fait sortir {sorted(fuite)}" if fuite else ""))
            continue
        # PARTIE PAR PARTIE, dans les deux sens : ce qui est admis dans l'onglet `fiche`
        # ne l'est pas dans l'onglet `activite`, même si les deux ont un écart.
        admis = ECARTS[cle].parties if cle in ECARTS else {}
        for nom, texte in sorted(parties.items()):
            fuite = _presents(interdits, texte)
            declare = set(admis.get(nom, ()))
            if fuite - declare:
                fautes.append(
                    f"{ou}, partie « {nom} » : {sorted(fuite - declare)} sort d'une "
                    "collection qu'on ne dépose pas, ou d'une charge retenue"
                    + (" — en plus de l'écart déclaré pour cette partie" if declare else ""))
            for s in sorted(declare - fuite):
                fautes.append(f"{ou}, partie « {nom} » : l'écart déclare « {s} », qui ne "
                              "sort plus — retirer la déclaration, elle rassure à tort")
        for nom in sorted(set(admis) - set(parties)):
            fautes.append(f"{ou} : l'écart vise la partie « {nom} », qui n'existe "
                          f"pas (parties : {sorted(parties)})")
        tout = "\n".join(parties.values())
        vu = _presents(propre, tout)
        du = {m for m in propre if m.startswith(attendu.porte)} if attendu.porte else set()
        if vu != du:
            fautes.append(
                f"{ou} : devait porter {sorted(du)} et porte {sorted(vu)} — "
                f"manque {sorted(du - vu)}, en trop {sorted(vu - du)}. Soit le filtre a "
                "changé, soit le LECTEUR ne voit plus le contenu (un classeur ou une "
                "archive lus en octets ne « portent » rien).")
        if attendu.temoin and attendu.temoin.lower() not in tout.lower():
            fautes.append(f"{ou} : le témoin « {attendu.temoin} » n'y est pas — "
                          "l'artefact est vide, ou il n'a pas été lu")
    return fautes


def _cles_attendues() -> set:
    """Les clés que le balayage DOIT avoir jouées, compte tenu d'un extra absent."""
    return {c for c in ATTENDUS
            if _xlsx_lisible() or c[2] not in ("xlsx", "--xlsx")}


# --------------------------------------------------------------------------- #
# Les cliquets de CONTENU
# --------------------------------------------------------------------------- #
def test_chaque_route_de_sortie_est_propre_ou_son_ecart_est_declare(client, seme,
                                                                    monkeypatch):
    """Les routes, voie ShareDocs comprise, chacune dans chaque format et chaque réglage.

    L'administrateur est le cas le plus exigeant, pour la raison écrite dans
    `test_depot_export` : sa portée est totale, donc aucune garde d'autorisation ne
    rattrape un périmètre mal passé.
    """
    vus = _balayer_routes(client, seme, seme["alpha"]["id"], monkeypatch)
    fautes = _juger(vus, seme["propre"], seme["interdits"])
    jouees = {c for c, _, _ in vus}
    for cle in sorted({c for c in _cles_attendues() if c[0] == "route"} - jouees):
        fautes.append(f"{cle} a un attendu et n'a pas été jouée")
    # La voie ShareDocs a bien joué CHAQUE export : sans ce compte, une boucle vide
    # passerait, et « le dépôt est propre » ne reposerait sur rien.
    deposes = {c for c, v, _ in vus if v.startswith("voie ShareDocs")}
    import routes.depot as depot
    dus = {("route", DEPOSER.replace("deposer", q), f)
           for q, fs in depot._FORMATS.items() for f in fs
           if f != "xlsx" or _xlsx_lisible()}
    assert deposes == dus, f"voie ShareDocs : joué {sorted(deposes)}, dû {sorted(dus)}"
    assert not fautes, ("Des routes font sortir ce qu'elles n'ont pas déclaré.\n  "
                        + "\n  ".join(fautes))


def test_chaque_outil_est_propre_ou_son_ecart_est_declare(seme, tmp_path):
    """Les outils hors application, lancés comme au shell avec `--collection`."""
    vus = _balayer_outils(seme, tmp_path)
    fautes = _juger(vus, seme["propre"], seme["interdits"])
    jouees = {c for c, _, _ in vus}
    for cle in sorted({c for c in _cles_attendues() if c[0] == "outil"} - jouees):
        fautes.append(f"{cle} a un attendu et n'a pas été jouée")
    assert not fautes, ("Des outils font sortir ce qu'ils n'ont pas déclaré.\n  "
                        + "\n  ".join(fautes))


# --------------------------------------------------------------------------- #
# Le semis — le mode d'échec d'un cliquet
# --------------------------------------------------------------------------- #
def test_le_semis_est_visible(client, seme):
    """Chaque sentinelle est RÉELLEMENT là, et un lecteur la voit quelque part.

    Les cliquets ci-dessus cherchent les termes de Bravo et les charges par leur ABSENCE :
    un décor qui ne les poserait pas les rendrait tous verts. Trois preuves, donc.

    1. EN BASE, relu et non déduit du décor : chaque terme existe avec la portée qu'on lui
       prête, les liaisons qui croisent la frontière existent, et le journal porte des
       charges — les marqueurs (que `_semer_une_charge_par_cible` relit déjà) ET une
       charge RÉELLE, écrite par l'API, qui nomme le tag de Bravo.
    2. À LA SORTIE, là où c'est permis : l'export de BRAVO porte les termes de Bravo —
       sans quoi un filtre qui tairait tout terme local passerait pour propre.
    3. PAR LE LECTEUR : la sauvegarde est la base entière, dépliée par le même lecteur
       que les artefacts ; toute sentinelle interdite doit y être vue. C'est la preuve
       qu'« absent » ci-dessus veut dire « filtré », pas « illisible ».
    """
    conn = sqlite3.connect(seme["db"])
    conn.row_factory = sqlite3.Row
    try:
        portee = {}
        for table, col in (("domaine", "nom"), ("attribut_dimension", "nom"),
                           ("attribut_valeur", "valeur"), ("tags", "label")):
            for r in conn.execute(f"SELECT {col} AS mot, collection_id FROM {table}"):
                portee[r["mot"]] = r["collection_id"]
        attendue = {**{m: None for m in seme["global"]},
                    **{m: seme["alpha"]["id"] for m in seme["a"]},
                    **{m: seme["bravo"]["id"] for m in seme["b"]}}
        # La seule sentinelle de Bravo qui soit GLOBALE : elle vit sous un axe de Bravo.
        attendue["valg-sous-bravo-9312"] = None
        for mot, cid in sorted(attendue.items()):
            assert mot in portee, f"« {mot} » n'est dans aucune table de vocabulaire"
            assert portee[mot] == cid, f"« {mot} » : portée {portee[mot]}, attendu {cid}"

        poses = {r[0] for r in conn.execute(
            "SELECT t.label FROM annotation_tags at JOIN tags t ON t.id = at.tag_id "
            "JOIN annotations a ON a.id = at.annotation_id WHERE a.region_id = ?",
            (seme["region_id"],))}
        assert {"tag-bravo-9402", "tag-alpha-9401", "tag-global-9400"} <= poses, poses
        situes = {r[0] for r in conn.execute(
            "SELECT v.valeur FROM region_attribut ra JOIN attribut_valeur v "
            "ON v.id = ra.valeur_id WHERE ra.region_id = ?", (seme["region_id"],))}
        assert {"val-bravo-9302", "val-alpha-9301"} <= situes, situes
        profil = {r[0] for r in conn.execute(
            "SELECT v.valeur FROM personnage_attribut pa JOIN attribut_valeur v "
            "ON v.id = pa.valeur_id")}
        assert {"valp-bravo-9311", "valp-global-9310"} <= profil, profil
        # L'axe GLOBAL rattaché au domaine de Bravo — l'état d'une base antérieure à
        # v24, que le décor pose par SQL. C'est lui qui éprouve que le nom du domaine ne
        # ressort pas en libellé d'un axe parfaitement lisible ; neutralisé, tout restait
        # vert (relecture du 2026-10-07).
        rattache = conn.execute(
            "SELECT d.nom FROM attribut_dimension a JOIN domaine d ON d.id = a.domaine_id "
            "WHERE a.nom = 'dim-global-9200'").fetchone()
        assert rattache and rattache[0] == "dom-bravo-9102", (
            "`dim-global-9200` n'est plus rattachée à `dom-bravo-9102` : la liaison "
            "axe global → domaine d'ailleurs n'est plus dans le décor")
        # Le contenu de l'album de Bravo, que `seme` pose par l'API.
        bulle = conn.execute(
            "SELECT r.ocr_texte, a.titre, n.note FROM regions r "
            "JOIN planches p ON p.id = r.planche_id JOIN albums a ON a.id = p.album_id "
            "LEFT JOIN annotations n ON n.region_id = r.id WHERE r.id = ?",
            (seme["region_bravo"],)).fetchone()
        assert tuple(bulle) == (TEXTE_AILLEURS, ALBUM_AILLEURS, NOTE_AILLEURS), tuple(bulle)
        poses_b = {r[0] for r in conn.execute(
            "SELECT t.label FROM annotation_tags at JOIN tags t ON t.id = at.tag_id "
            "JOIN annotations a ON a.id = at.annotation_id WHERE a.region_id = ?",
            (seme["region_bravo"],))}
        assert poses_b == {"tag-bravo-9402"}, poses_b

        reelles = [r[0] for r in conn.execute(
            "SELECT apres FROM evenement WHERE agent = 'decor' AND apres LIKE "
            "'%tag-bravo-9402%'")]
        assert reelles, ("aucune charge RÉELLE ne nomme le tag de Bravo : le journal du "
                         "décor ne prouverait plus que taire les charges sert à quelque "
                         "chose")
        assert conn.execute("SELECT COUNT(*) FROM activite WHERE portee LIKE ?",
                            (f"%{PORTEE_AILLEURS}%",)).fetchone()[0] == 1
    finally:
        conn.close()

    # (2) Bravo exporte Bravo — chaque route, dans la famille qu'elle annonce.
    bravo = _balayer_routes(client, seme, seme["bravo"]["id"], album=False)
    assert {c for c, _, _ in bravo} == {c for c in _cles_attendues()
                                        if c[0] == "route" and DEPOT in c[1]}
    verbatim_vu = False
    for cle, variante, parties in bravo:
        tout = "\n".join(parties.values())
        dus = {m for m in seme["b"] if m.startswith(ATTENDUS[cle].porte)} \
            if ATTENDUS[cle].porte else set()
        manque = dus - _presents(dus, tout)
        assert not manque, (f"{cle[1]} [{cle[2]}] ({variante}) : Bravo n'exporte plus "
                            f"son propre {sorted(manque)}")
        # Le périmètre d'ALBUMS, à l'endroit : là où Alpha nomme son album, Bravo nomme
        # le sien — sans quoi un titre absent de partout passerait pour « bien filtré ».
        if ATTENDUS[cle].porte in (TOUT, TAGS):
            assert ALBUM_AILLEURS in tout and NOTE_AILLEURS in tout, (
                f"{cle[1]} [{cle[2]}] ({variante}) : l'album de Bravo ou sa note manque "
                "de l'export de Bravo")
            if '"verbatim": true' in variante:
                assert TEXTE_AILLEURS in tout, (
                    f"{cle[1]} [{cle[2]}] ({variante}) : `verbatim` ne fait pas sortir "
                    "le texte de la bulle de Bravo — la sentinelle de texte ne mesure rien")
                verbatim_vu = True
    assert verbatim_vu, "aucune variante `verbatim` jouée pour Bravo"

    # (3) Le lecteur voit TOUT dans la sauvegarde — lue en OCTETS, c'est un fichier SQLite.
    rep = client.get("/api/sauvegarde", headers=ADMIN)
    assert rep.status_code == 200, rep.text
    base = _brut(rep.content)
    invisibles = set(seme["interdits"]) - _presents(seme["interdits"], base)
    assert not invisibles, (
        f"sentinelles introuvables dans la base entière : {sorted(invisibles)} — le "
        "décor ne les pose pas, ou le lecteur ne sait pas lire. Un cliquet qui cherche "
        "une absence est alors vert pour la mauvaise raison.")


def test_le_lecteur_ouvre_un_classeur(seme, client):
    """Le lecteur de XLSX rend des FEUILLES nommées, et leur contenu.

    Un écart se déclare par feuille (`fiche`, `activite`) : si le lecteur rendait le
    classeur en un bloc, ou ses octets, les déclarations ne viseraient plus rien — et
    c'est par l'égalité du vocabulaire qu'on le verrait, pas ici. Ce test dit la chose
    directement, sur l'artefact réel.
    """
    if not _xlsx_lisible():
        pytest.skip("openpyxl absent (extra d'export)")
    rep = client.get(_META.replace("{collection_id}", str(seme["alpha"]["id"])),
                     params={"format": "xlsx"}, headers=ADMIN)
    assert rep.status_code == 200, rep.text
    assert not _presents(seme["propre"], rep.content.decode("latin-1")), (
        "les octets BRUTS du classeur portent une sentinelle en clair : ce test ne "
        "prouve plus que la lecture en octets est vacante")
    feuilles = _parties(rep.content)
    assert {"fiche", "arbre", "vocabulaire", "tags", "activite", "evenement"} <= set(feuilles)
    assert "tag-alpha-9401" in feuilles["tags"]
    assert "dim-alpha-9201" in feuilles["vocabulaire"]


def test_le_balayage_dit_ce_qu_il_n_a_pas_vu(capsys):
    """N'échoue pas : AFFICHE ce qu'un extra absent empêche de jouer (patron d'AUTH-5).

    Sans `openpyxl`, le classeur n'est ni produit (503 nommé) ni lisible : trois sorties
    ne sont regardées par personne jusqu'à ce que l'image tourne. Un cliquet partiel qui
    se tait rassure à tort.
    """
    if _xlsx_lisible():
        return
    with capsys.disabled():
        print("\n  AUTH-11 — balayage PARTIEL : `openpyxl` absent, le classeur XLSX "
              "n'est joué ni par la route, ni par la voie ShareDocs, ni par l'outil "
              "(pip install -r requirements-export.txt)")
